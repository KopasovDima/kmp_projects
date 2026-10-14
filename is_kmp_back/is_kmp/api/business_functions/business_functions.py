import zipfile
from collections import defaultdict
from http import HTTPStatus
from io import BytesIO

from django.contrib.postgres.aggregates import StringAgg
from django.db.models import Q, CharField, F
from django.db.models.functions import Cast
from django.http import FileResponse
from django.views.decorators.http import require_GET, require_POST
from openpyxl.styles import Alignment, PatternFill, Font
from openpyxl.workbook import Workbook

from is_kmp.consts.roles import ROLE_IT_SPECIALIST, ROLE_ADMIN, ROLE_BF_MANAGER
from is_kmp.exceptions.bad_request_exception import BadRequestException
from is_kmp.exceptions.forbidden_exception import ForbiddenException
from is_kmp.exceptions.not_found_exception import NotFoundException
from is_kmp.models.business_function import BusinessFunction
from is_kmp.models.business_function_by_company import BusinessFunctionByCompany
from is_kmp.models.business_function_history import BusinessFunctionHistory
from is_kmp.models.company import Company
from is_kmp.models.company_software import CompanySoftware
from is_kmp.models.implementation_status import ImplementationStatus
from is_kmp.models.software import Software
from is_kmp.services.business_functions.business_function_by_company import \
    get_business_function_by_company_by_id_or_exception
from is_kmp.services.business_functions.business_functions import get_business_function_by_id_or_exception
from is_kmp.services.companies.companies import get_company_with_all_children
from is_kmp.services.software.company_software import get_company_software_by_id_or_exception
from is_kmp.services.software.software import get_software_by_id_or_exception
from is_kmp.utils.decorators import require_auth, require_PATCH, require_PUT, require_DELETE, require_role
from is_kmp.utils.permissions import check_have_access_to_company
from is_kmp.utils.request_params import get_request_value_str, get_request_value_int, get_request_value_bool, \
    get_patch_request_data


@require_GET
@require_auth
def get_business_functions(request, _, user_company_id, __):
    include_software = get_request_value_str(request.GET, 'include_software')
    root_business_function_id = get_request_value_int(request.GET, 'root_business_function_id')

    business_functions = BusinessFunction.objects.filter(is_deleted=False).select_related('parent')

    if root_business_function_id:
        exclude_level_0_ids = [
            level_0_id.get('id') for level_0_id in BusinessFunction.objects.filter(level=0).values('id')
            if level_0_id.get('id') != root_business_function_id
        ]
        business_functions = business_functions.exclude(pk__in=exclude_level_0_ids)

    if include_software:
        business_functions = (
            business_functions
            .annotate(software_name=StringAgg(
                'business_function_by_company__company_software__software__name',
                ', ',
                ordering=['business_function_by_company__company_software__software__name'],
                filter=Q(business_function_by_company__company_id=user_company_id),
            ))
            .annotate(value=StringAgg(
                Cast('business_function_by_company__value', output_field=CharField()),
                ', ',
                ordering=['business_function_by_company__value'],
                filter=Q(business_function_by_company__company_id=user_company_id),
            ))
            .order_by('sorting_order', 'id')
        )
    else:
        business_functions = business_functions.order_by('id')

    children_map = defaultdict(list)
    software_map = dict()

    for business_function in business_functions:
        if include_software:
            software_map[business_function.id] = (
                business_function.software_name,
                business_function.value,
            )
        if business_function.parent_id:
            children_map[business_function.parent_id].append(business_function)

    def build_tree(bf):
        # После ручного импорта данных появляются пробелы в конце имени
        # if bf.name != bf.name.strip():
        #     bf.name = bf.name.strip()
        #     bf.save()

        bf_info = {
            'id': bf.id,
            'name': bf.name,
            'level': bf.level,
            'last_updated': bf.last_update,
            'sorting_order': bf.sorting_order,
        }

        if bf.level != 3:
            bf_info['children'] = [build_tree(child) for child in children_map.get(bf.id, [])]

        if software_info := software_map.get(bf.id):
            if bf.level in [2, 3]:
                value = software_info[1]
                bf_info['value'] = 'true' in value if value else None
            if bf_info.get('value') and (software := software_info[0]):
                bf_info['software'] = software

        return bf_info

    root_business_functions = [
        business_function for business_function in business_functions if business_function.parent_id is None
    ]

    return {
        'business_functions': [build_tree(business_function) for business_function in root_business_functions],
    }


@require_POST
@require_role(ROLE_IT_SPECIALIST)
def add_business_function(request, user_id, _, __):
    name = get_request_value_str(request.POST, 'name', True)
    parent_id = get_request_value_int(request.POST, 'parent_id', True)

    parent_business_function = get_business_function_by_id_or_exception(parent_id)

    business_function = BusinessFunction()

    business_function.name = name
    business_function.parent = parent_business_function
    business_function.level = parent_business_function.level + 1

    business_function.ftr_example = ''
    business_function.priority_example = None
    business_function.reason_example = ''

    business_function.save()

    business_function.sorting_order = business_function.pk

    business_function.save()

    business_function_history = BusinessFunctionHistory()

    business_function_history.business_function_id = business_function.pk
    business_function_history.action_type = 'created'
    business_function_history.changed_by_user = user_id
    business_function_history.changes = (
        f'name:{name}||parent_id:{parent_id}||'
        f'level:{business_function.level}||sorting_order={business_function.sorting_order}'
    )

    business_function_history.save()

    return {
        'created_business_function': {
            'id': business_function.pk,
            'name': business_function.name,
            'level': business_function.level,
            'last_updated': business_function.last_update,
        }
    }


@require_PUT
@require_role(ROLE_IT_SPECIALIST)
def update_business_function(request, user_id, _, __, business_function_id: int):
    request_data = get_patch_request_data(request)

    name = get_request_value_str(request_data, 'name', True)

    business_function = get_business_function_by_id_or_exception(business_function_id)

    if business_function.is_deleted:
        raise ForbiddenException('Объект удалён и недоступен для редактирования')

    business_function.name = name

    business_function.save()

    business_function_history = BusinessFunctionHistory()

    business_function_history.business_function_id = business_function.pk
    business_function_history.action_type = 'updated'
    business_function_history.changed_by_user = user_id
    business_function_history.changes = f'name:{name}'

    business_function_history.save()

    return {
        'updated_business_function': {
            'id': business_function.id,
            'name': business_function.name,
            'level': business_function.level,
            'last_updated': business_function.last_update,
        }
    }


@require_PUT
@require_role(ROLE_BF_MANAGER)
def update_business_function_examples(request, _, __, ___, business_function_id: int):
    request_data = get_patch_request_data(request)

    ftr_example = get_request_value_str(request_data, 'ftr_example') or ''
    priority_example = get_request_value_str(request_data, 'priority_example') or ''
    reason_example = get_request_value_str(request_data, 'reason_example') or ''

    business_function = get_business_function_by_id_or_exception(business_function_id)

    if business_function.is_deleted:
        raise ForbiddenException('Объект удалён и недоступен для редактирования')

    if business_function.level not in [2, 3]:
        raise ForbiddenException('Пример заполнения доступен только по БФ 2 или 3 уровня')

    business_function.ftr_example = ftr_example
    business_function.priority_example = priority_example
    business_function.reason_example = reason_example

    business_function.save()

    return 'OK'


@require_DELETE
@require_role(ROLE_ADMIN)
def delete_business_function(_, user_id, __, ___, business_function_id: int):
    business_function = get_business_function_by_id_or_exception(business_function_id)

    if business_function.is_deleted:
        raise ForbiddenException('Объект уже удалён')

    business_function.is_deleted = True
    business_function.save()

    business_function_history = BusinessFunctionHistory()

    business_function_history.business_function_id = business_function_id
    business_function_history.action_type = 'deleted'
    business_function_history.changed_by_user = user_id
    business_function_history.changes = ''

    business_function_history.save()

    return 'ОК'


@require_GET
@require_auth
def get_business_function_by_company_details(_, __, user_company_id, ___, business_function_id: int):
    business_function = get_business_function_by_id_or_exception(business_function_id)

    if business_function.is_deleted:
        raise ForbiddenException('Объект удалён и недоступен')

    if business_function.level not in [2, 3]:
        raise ForbiddenException('Детальная информация доступна только по БФ 2 или 3 уровня')

    business_function_by_company = (
        business_function.business_function_by_company
        .filter(company_id=user_company_id).first()
    )

    if not business_function_by_company:
        business_function_by_company = BusinessFunctionByCompany()

        business_function_by_company.company_id = user_company_id
        business_function_by_company.business_function = business_function

        business_function_by_company.save()

    return {
        'business_function_by_company': {
            'id': business_function_by_company.pk,
            'value': business_function_by_company.value,
            'software_usage_status': business_function_by_company.software_usage_status or '',
            'comment': business_function_by_company.comment or '',
            'reason': business_function_by_company.reason or '',
            'priority': business_function_by_company.priority or '',
            'ftr': business_function_by_company.ftr or '',
            'ru_software_availability_status': business_function_by_company.ru_software_availability_status or '',
            'target_software': {
                'id': business_function_by_company.target_software_id,
                'name': business_function_by_company.target_software.name,
            } if business_function_by_company.target_software else None,
            'company_software': [
                {
                    'id': company_software.pk,
                    'software': {
                        'id': company_software.software.pk,
                        'name': company_software.software.name,
                    },
                } for company_software in business_function_by_company.company_software.order_by('software__name')
            ],
            'examples': {
                'ftr_example': business_function.ftr_example or '',
                'priority_example': business_function.priority_example or '',
                'reason_example': business_function.reason_example or '',
            },
        }
    }


@require_PUT
@require_role(ROLE_IT_SPECIALIST)
def update_business_function_by_company(request, _, user_company_id, user_roles, business_function_by_company_id: int):
    request_data = get_patch_request_data(request)

    value = get_request_value_bool(request_data, 'value') or False
    reason = get_request_value_str(request_data, 'reason') or ''
    priority = get_request_value_str(request_data, 'priority') or ''
    ftr = get_request_value_str(request_data, 'ftr') or ''
    ru_software_availability_status = get_request_value_str(request_data, 'ru_software_availability_status') or ''
    software_usage_status = get_request_value_str(request_data, 'software_usage_status') or ''
    target_software_id = get_request_value_int(request_data, 'target_software_id')
    comment = get_request_value_str(request_data, 'comment') or ''

    if ru_software_availability_status and ru_software_availability_status not in ['0', '1', '2', '3']:
        raise BadRequestException(f'Недопустимое значение ru_software_availability_status: {ru_software_availability_status}')

    if software_usage_status and software_usage_status not in [
        'Используется без ПО', 'Используется с российским ПО', 'Используется с иностранным ПО',
    ]:
        raise BadRequestException(f'Недопустимое значение software_usage_status: {software_usage_status}')

    business_function_by_company = get_business_function_by_company_by_id_or_exception(business_function_by_company_id)

    if business_function_by_company.business_function.is_deleted:
        raise ForbiddenException('Объект удалён и недоступен для редактирования')

    if target_software_id:
        target_software = get_software_by_id_or_exception(target_software_id)
    else:
        target_software = None

    if target_software and not target_software.register_number:
        raise BadRequestException('В качестве целевого ПО можно выбрать только ПО, состоящее в реестре российского ПО')

    check_have_access_to_company(user_company_id, business_function_by_company.company_id, user_roles)

    if target_software_id != business_function_by_company.target_software_id:
        business_function_by_company.target_software = target_software

    business_function_by_company.value = value

    business_function_by_company.reason = reason

    business_function_by_company.ru_software_availability_status = ru_software_availability_status

    business_function_by_company.software_usage_status = software_usage_status
    business_function_by_company.comment = comment
    business_function_by_company.priority = priority
    business_function_by_company.ftr = ftr

    business_function_by_company.save()

    return {
        'updated_business_function_by_company': {
            'id': business_function_by_company.pk,
            'value': business_function_by_company.value,
            'reason': business_function_by_company.reason or '',
            'priority': business_function_by_company.priority or '',
            'ftr': business_function_by_company.ftr or '',
            'ru_software_availability_status': business_function_by_company.ru_software_availability_status or '',
            'software_usage_status': business_function_by_company.software_usage_status or '',
            'target_software': {
                'id': business_function_by_company.target_software_id,
                'name': business_function_by_company.target_software.name,
            } if business_function_by_company.target_software else None,
            'comment': business_function_by_company.comment or '',
        }
    }


@require_PATCH
@require_role(ROLE_IT_SPECIALIST)
def attach_company_software_to_business_function(
        _, __, user_company_id, ___, business_function_id: int, company_software_id: int,
):
    business_function = get_business_function_by_id_or_exception(business_function_id)

    if business_function.is_deleted:
        raise ForbiddenException('Объект удалён и недоступен для редактирования')

    if business_function.level not in [2, 3]:
        raise ForbiddenException('Редактирование доступно только для БФ 2 или 3 уровня')

    company_software = get_company_software_by_id_or_exception(company_software_id)

    if user_company_id != company_software.company_id:
        raise ForbiddenException('Невозможно прикрепить ПО другой компании')

    business_function_by_company = business_function.business_function_by_company.filter(
        company_id=user_company_id
    ).first()

    if not business_function_by_company:
        raise BadRequestException('Не найдена подходящая модель БФ компании')

    existing_company_software_in_business_function = business_function_by_company.company_software.filter(
        pk=company_software_id
    ).first()

    if existing_company_software_in_business_function:
        raise ForbiddenException('Данное ПО уже прикреплено к бизнес функции')

    business_function_by_company.company_software.add(company_software)

    business_function_by_company.value = True

    business_function_by_company.save()

    return {
        'business_function_by_company': {
            'company_software': [
                {
                    'id': company_software.pk,
                    'software': {
                        'id': company_software.software.pk,
                        'name': company_software.software.name,
                    },
                } for company_software in business_function_by_company.company_software.order_by('software__name')
            ],
        }
    }


@require_PATCH
@require_role(ROLE_IT_SPECIALIST)
def attach_new_company_software_to_business_function(
        _, __, user_company_id, ___, business_function_id: int, software_id: int,
):
    business_function = get_business_function_by_id_or_exception(business_function_id)

    if business_function.is_deleted:
        raise ForbiddenException('Объект удалён и недоступен для редактирования')

    if business_function.level not in [2, 3]:
        raise ForbiddenException('Редактирование доступно только для БФ 2 или 3 уровня')

    business_function_by_company = business_function.business_function_by_company.filter(
        company_id=user_company_id
    ).first()

    if not business_function_by_company:
        raise BadRequestException('Не найдена подходящая модель БФ компании')

    existing_company_software_in_business_function = business_function_by_company.company_software.filter(
        software_id=software_id
    ).first()

    if existing_company_software_in_business_function:
        raise ForbiddenException('Данное ПО уже прикреплено к бизнес функции')

    get_software_by_id_or_exception(software_id)

    first_implementation_status = ImplementationStatus.objects.order_by('ordering_number', 'id').first()

    if not first_implementation_status:
        raise NotFoundException('Отсутствуют статусы реализации')

    company_software = CompanySoftware()

    company_software.company_id = user_company_id
    company_software.software_id = software_id
    company_software.copies_number = 0
    company_software.additionally_needed_copies_number = 0
    company_software.implementation_status = first_implementation_status

    company_software.save()

    business_function_by_company.company_software.add(company_software)

    business_function_by_company.value = True

    business_function_by_company.save()

    return {
        'business_function_by_company': {
            'company_software': [
                {
                    'id': company_software.pk,
                    'software': {
                        'id': company_software.software.pk,
                        'name': company_software.software.name,
                    },
                } for company_software in business_function_by_company.company_software.order_by('software__name')
            ],
        }
    }


@require_PATCH
@require_role(ROLE_IT_SPECIALIST)
def detach_company_software_from_business_function(
        request, _, user_company_id, __, business_function_id: int, company_software_id: int,
):
    request_data = get_patch_request_data(request)

    keep_value_true = get_request_value_bool(request_data, 'keep_value_true') or False

    business_function = get_business_function_by_id_or_exception(business_function_id)

    if business_function.is_deleted:
        raise ForbiddenException('Объект удалён и недоступен для редактирования')

    if business_function.level not in [2, 3]:
        raise ForbiddenException('Редактирование доступно только для БФ 2 или 3 уровня')

    company_software = get_company_software_by_id_or_exception(company_software_id)

    if user_company_id != company_software.company_id:
        raise ForbiddenException('Невозможно прикрепить ПО другой компании')

    business_function_by_company = business_function.business_function_by_company.filter(
        company_id=user_company_id
    ).first()

    if not business_function_by_company:
        raise BadRequestException('Не найдена подходящая модель БФ компании')

    existing_company_software_in_business_function = business_function_by_company.company_software.filter(
        pk=company_software_id
    ).first()

    if not existing_company_software_in_business_function:
        raise ForbiddenException('ПО не прикреплено к бизнес функции')

    business_function_by_company.company_software.remove(company_software)

    if keep_value_true or (
            business_function_by_company.reason or
            business_function_by_company.priority or
            business_function_by_company.ftr or
            business_function_by_company.ru_software_availability_status or
            business_function_by_company.software_usage_status or
            business_function_by_company.comment or
            business_function_by_company.target_software_id or
            business_function_by_company.company_software.count()
    ):
        business_function_by_company.value = True
    else:
        business_function_by_company.value = False

    business_function_by_company.save()

    return {
        'business_function_by_company': {
            'company_software': [
                {
                    'id': company_software.pk,
                    'software': {
                        'id': company_software.software.pk,
                        'name': company_software.software.name,
                    },
                } for company_software in business_function_by_company.company_software.order_by('software__name')
            ],
        }
    }


@require_PATCH
@require_role(ROLE_IT_SPECIALIST)
def change_business_function_order(request, user_id, _, __, business_function_id: int):
    request_data = get_patch_request_data(request)

    business_function = get_business_function_by_id_or_exception(business_function_id)

    if business_function.is_deleted:
        raise ForbiddenException('Объект удалён и недоступен для редактирования')

    new_position = get_request_value_int(request_data, 'new_position', True)
    new_parent_id = get_request_value_int(request_data, 'new_parent_id')

    if new_parent_id:
        new_parent_business_function = get_business_function_by_id_or_exception(new_parent_id)

        if new_parent_business_function.level != business_function.parent.level:
            raise BadRequestException('Нельзя сменить уровень БФ при изменения сортировки')

        business_function.parent = new_parent_business_function

    business_function.sorting_order = new_position

    business_function.save()

    other_business_functions = (
        BusinessFunction.objects
        .exclude(pk=business_function_id)
        .filter(parent_id=business_function.parent_id, sorting_order__gte=new_position)
        .order_by('sorting_order', 'id')
    )

    history_objects = []

    business_function_history = BusinessFunctionHistory()

    business_function_history.business_function_id = business_function.pk
    business_function_history.action_type = 'changed_order'
    business_function_history.changed_by_user = user_id
    business_function_history.changes = (
        f'{("parent_id:" + str(new_parent_id) + "||") if new_parent_id else ""}sorting_order={new_position}'
    )
    history_objects.append(business_function_history)

    for bf in other_business_functions:
        new_position += 1
        bf.sorting_order = new_position
        bf.save()

        business_function_history = BusinessFunctionHistory()

        business_function_history.business_function_id = bf.pk
        business_function_history.action_type = 'changed_order'
        business_function_history.changed_by_user = user_id
        business_function_history.changes = (
            f'sorting_order={new_position}'
        )

        history_objects.append(business_function_history)

    BusinessFunctionHistory.objects.bulk_create(history_objects)

    return 'OK'


@require_PATCH
@require_role(ROLE_IT_SPECIALIST)
def reset_company_business_function(request, _, user_company_id, user_roles, business_function_by_company_id: int):
    request_data = get_patch_request_data(request)

    business_function_by_company = get_business_function_by_company_by_id_or_exception(business_function_by_company_id)

    keep_old_values = get_request_value_bool(request_data, 'keep_old_values')

    check_have_access_to_company(user_company_id, business_function_by_company.company_id, user_roles)

    if business_function_by_company.business_function.is_deleted:
        raise ForbiddenException('Объект удалён и недоступен для редактирования')

    business_function_by_company.value = None

    if not keep_old_values:
        business_function_by_company.reason = ''

        business_function_by_company.ru_software_availability_status = ''

        business_function_by_company.software_usage_status = ''
        business_function_by_company.comment = ''
        business_function_by_company.target_software = None
        business_function_by_company.priority = ''
        business_function_by_company.ftr = ''

        business_function_by_company.company_software.clear()

    business_function_by_company.save()

    return {
        'updated_business_function_by_company': {
            'id': business_function_by_company.pk,
            'value': business_function_by_company.value,
            'reason': business_function_by_company.reason or '',
            'priority': business_function_by_company.priority or '',
            'ftr': business_function_by_company.ftr or '',
            'ru_software_availability_status': business_function_by_company.ru_software_availability_status or '',
            'software_usage_status': business_function_by_company.software_usage_status or '',
            'target_software': {
                'id': business_function_by_company.target_software_id,
                'name': business_function_by_company.target_software.name,
            } if business_function_by_company.target_software else None,
            'comment': business_function_by_company.comment or '',
            'company_software': [
                {
                    'id': company_software.pk,
                    'software': {
                        'id': company_software.software.pk,
                        'name': company_software.software.name,
                    },
                } for company_software in business_function_by_company.company_software.order_by('software__name')
            ],
        }
    }


@require_GET
@require_auth
def download_report(_, __, ___, ____):
    root_business_functions = (
        BusinessFunction.objects
        .filter(parent_id__isnull=True, is_deleted=False)
        .order_by('sorting_order', 'id')
    )

    business_functions = (
        BusinessFunction.objects
        .filter(is_deleted=False)
        .select_related('parent')
        .order_by('sorting_order', 'id')
    )

    cell_alignment = Alignment(wrapText=True, horizontal='left', vertical='center')
    cell_fill = PatternFill(
        start_color='9dc3e6',
        end_color='9dc3e6',
        fill_type='darkHorizontal',
    )

    def write_business_functions_to_excel_sheet(ws, number, position, items):
        for index, business_function in enumerate(items, start=1):
            position += 1

            number_cell = ws.cell(row=position, column=2, value=f'{number}{index}')
            number_cell.alignment = cell_alignment

            data_cell = ws.cell(row=position, column=3, value=business_function.get('name') or '---')
            data_cell.alignment = cell_alignment

            if business_function.get('level') in [1, 2]:
                number_cell.fill = cell_fill
                data_cell.fill = cell_fill

            if children := business_function.get('children'):
                position = write_business_functions_to_excel_sheet(ws, f'{number}{index}.', position, children)
            elif business_function.get('level') != 3:
                position += 1

        return position

    zip_buffer = BytesIO()

    with zipfile.ZipFile(zip_buffer, 'w', zipfile.ZIP_DEFLATED) as zip_file:
        for root_business_function in root_business_functions:
            wb = Workbook()

            children_map = defaultdict(list)

            for business_function in business_functions:
                if business_function.parent_id:
                    children_map[business_function.parent_id].append(business_function)

            def build_tree(bf):
                bf_info = {
                    'name': bf.name,
                    'level': bf.level,
                }

                if bf.level != 3:
                    bf_info['children'] = [build_tree(child) for child in children_map.get(bf.id, [])]

                return bf_info

            business_functions_tree = build_tree(root_business_function)

            ws = wb.active
            ws.title = business_functions_tree.get('name') or '---'

            position = 4

            ws.column_dimensions['C'].width = 80

            ws.cell(row=position, column=2, value='№').alignment = cell_alignment
            ws.cell(row=position, column=3, value='Бизнес-функции').alignment = cell_alignment

            if business_functions_tree.get('level') == 0:
                for index, business_function in enumerate(business_functions_tree.get('children', []), start=1):
                    position += 1
                    ws.cell(row=position, column=2, value=index).alignment = cell_alignment
                    ws.cell(
                        row=position, column=3, value=business_function.get('name') or '---'
                    ).alignment = cell_alignment

                for index, business_function in enumerate(business_functions_tree.get('children', []), start=1):
                    ws = wb.create_sheet(f'{index}. {business_function.get("name") or "---"} ')

                    position = 4

                    ws.column_dimensions['C'].width = 120

                    ws.cell(row=position, column=2, value='№').alignment = cell_alignment
                    ws.cell(row=position, column=3, value='Бизнес-функции').alignment = cell_alignment

                    write_business_functions_to_excel_sheet(
                        ws,
                        f'{index}.',
                        position,
                        business_function.get('children', []),
                    )
            else:
                write_business_functions_to_excel_sheet(
                    ws,
                    '',
                    position,
                    business_functions_tree.get('children', []),
                )

            excel_buffer = BytesIO()
            wb.save(excel_buffer)
            excel_buffer.seek(0)

            zip_file.writestr(f'{root_business_function.name}.xlsx', excel_buffer.getvalue())

    zip_buffer.seek(0)

    return FileResponse(
        zip_buffer,
        filename='БизнесФункции.zip',
        content_type='application/zip',
        status=HTTPStatus.OK,
        as_attachment=True,
    )


@require_GET
@require_auth
def download_full_report(request, _, user_company_id, __):
    consolidated = get_request_value_bool(request.GET, 'consolidated')
    allow_empty_values = get_request_value_bool(request.GET, 'allow_empty_values')

    root_business_functions = (
        BusinessFunction.objects
        .filter(parent_id__isnull=True, is_deleted=False)
        .order_by('sorting_order', 'id')
    )

    user_company = Company.objects.filter(id=user_company_id).first()

    companies_ids = [user_company_id]
    companies = [user_company]

    if consolidated:
        companies = get_company_with_all_children(user_company_id)

        if len(companies) > 1:
            companies_ids = [company.id for company in companies]
        else:
            consolidated = False

    def get_company_business_functions(company_id):
        return (
            BusinessFunction.objects
            .filter(is_deleted=False)
            .select_related('parent')
            .order_by('sorting_order', 'id')
            .annotate(software_id=StringAgg(
                Cast('business_function_by_company__company_software__software__id', output_field=CharField()),
                '$$$',
                ordering=['business_function_by_company__company_software__software__name'],
                filter=Q(business_function_by_company__company_id=company_id),
            ))
            .annotate(software_usage_status=StringAgg(
                'business_function_by_company__software_usage_status',
                '$$$',
                ordering=['business_function_by_company__software_usage_status'],
                filter=Q(business_function_by_company__company_id=company_id),
            ))
            .annotate(ru_software_availability_status=StringAgg(
                'business_function_by_company__ru_software_availability_status',
                '$$$',
                ordering=['business_function_by_company__ru_software_availability_status'],
                filter=Q(business_function_by_company__company_id=company_id),
            ))
            .annotate(ftr=StringAgg(
                'business_function_by_company__ftr',
                '$$$',
                ordering=['business_function_by_company__ftr'],
                filter=Q(business_function_by_company__company_id=company_id),
            ))
            .annotate(priority=StringAgg(
                'business_function_by_company__priority',
                '$$$',
                ordering=['business_function_by_company__priority'],
                filter=Q(business_function_by_company__company_id=company_id),
            ))
            .annotate(reason=StringAgg(
                'business_function_by_company__reason',
                '$$$',
                ordering=['business_function_by_company__reason'],
                filter=Q(business_function_by_company__company_id=company_id),
            ))
            .annotate(comment=StringAgg(
                'business_function_by_company__comment',
                '$$$',
                ordering=['business_function_by_company__comment'],
                filter=Q(business_function_by_company__company_id=company_id),
            ))
            .annotate(target_software_id=StringAgg(
                Cast('business_function_by_company__target_software_id', output_field=CharField()),
                '$$$',
                ordering=['business_function_by_company__target_software_id'],
                filter=Q(business_function_by_company__company_id=company_id),
            ))
            .annotate(value=StringAgg(
                Cast('business_function_by_company__value', output_field=CharField()),
                '$$$',
                ordering=['business_function_by_company__value'],
                filter=Q(business_function_by_company__company_id=company_id),
            ))
            .order_by('sorting_order', 'id')
        )

    companies_business_functions = dict()

    for company_id in companies_ids:
        companies_business_functions[company_id] = get_company_business_functions(company_id)

    cell_alignment = Alignment(wrapText=True, horizontal='center', vertical='center')

    children_map = defaultdict(list)
    company_details_map = dict()

    def normalize_value(value):
        if not value:
            return ''
        return ', '.join(list(set(value.split('$$$'))))

    all_software_ids = set()

    for company_id in companies_ids:
        details_map = dict()
        company_details_map[company_id] = details_map

        for business_function in companies_business_functions[company_id]:
            software_ids = []
            if business_function.software_id:
                for software_id in list(set(business_function.software_id.split('$$$'))):
                    software_id = int(software_id)
                    if software_id not in software_ids:
                        software_ids.append(software_id)
                all_software_ids = all_software_ids.union(set(software_ids))

            details_map[business_function.id] = {
                'software_ids': list(software_ids),
                'software_usage_status': normalize_value(business_function.software_usage_status),
                'ru_software_availability_status': normalize_value(business_function.ru_software_availability_status),
                'ftr': normalize_value(business_function.ftr),
                'priority': normalize_value(business_function.priority),
                'reason': normalize_value(business_function.reason),
                'target_software_id': normalize_value(business_function.target_software_id),
                'comment': normalize_value(business_function.comment),
            }

            if details_map[business_function.id]['target_software_id']:
                details_map[business_function.id]['target_software_id'] = int(
                    details_map[business_function.id]['target_software_id']
                )
                all_software_ids = all_software_ids.union({details_map[business_function.id]['target_software_id']})

            if business_function.value and 'true' in business_function.value:
                details_map[business_function.id]['value'] = not not (
                        details_map[business_function.id]['software_ids'] or
                        details_map[business_function.id]['software_usage_status'] or
                        details_map[business_function.id]['ru_software_availability_status'] or
                        details_map[business_function.id]['ftr'] or
                        details_map[business_function.id]['priority'] or
                        details_map[business_function.id]['reason'] or
                        details_map[business_function.id]['target_software_id'] or
                        details_map[business_function.id]['comment']
                ) or allow_empty_values
            else:
                details_map[business_function.id]['value'] = False

            if ru_software_availability_status := details_map[business_function.id]['ru_software_availability_status']:
                ru_software_availability_status_values = {
                    '0': 'Нет решений',
                    '1': 'Необходима серьезная доработка решений',
                    '2': 'Необходима доработка решений',
                    '3': 'Российские решения закрывают критически важный функционал',
                }

                details_map[business_function.id]['ru_software_availability_status'] = (
                    ru_software_availability_status_values[str(ru_software_availability_status)]
                )

            if user_company_id == company_id and business_function.parent_id:
                children_map[business_function.parent_id].append(business_function)

    def build_tree(bf):
        bf_info = {
            'id': bf.id,
            'name': bf.name,
            'level': bf.level,
        }

        if bf.ftr_example:
            bf_info['ftr_example'] = bf.ftr_example

        if bf.priority_example:
            bf_info['priority_example'] = bf.priority_example

        if bf.reason_example:
            bf_info['reason_example'] = bf.reason_example

        if bf.level != 3:
            bf_info['children'] = [build_tree(child) for child in children_map.get(bf.id, [])]

        return bf_info

    all_software = (
        Software.objects
        .filter(id__in=list(all_software_ids))
        .annotate(main_class_name=F('main_class__name'))
        .annotate(additional_classes_names=StringAgg(
            'additional_classes__name',
            ', ',
            ordering=['additional_classes__number'],
        ))
    )

    software_dict = {}

    for software in all_software:
        software_dict[software.id] = {
            'name': f'{software.name} ({software.register_number})' if software.register_number else software.name,
            'main_class_name': software.main_class_name,
            'additional_classes_names': software.additional_classes_names or ''
        }

    wb = Workbook()

    ws = wb.active
    ws.title = '0'

    position = 2

    columns_width = {
        'bf_name': 20,
        'bf_ftr_example': 20,
        'bf_priority_example': 20,
        'bf_reason_example': 20,
        'company_name': 25,
        'ftr': 20,
        'priority': 15,
        'reason': 20,
        'ru_software_availability_status': 24,
        'software_usage_status': 22,
        'target_software_names': 20,
        'target_software_main_classes': 20,
        'target_software_additional_classes': 20,
        'software_in_use_names': 20,
        'comment': 20,
    }

    def insert_value_into_cell(row, column, value, is_bold=False):
        cell = ws.cell(row=row, column=column, value=value)
        if is_bold:
            cell.font = Font(bold=True)
        cell.alignment = cell_alignment

    def insert_row_data(row_pos, bf_id, values):
        company_name_values = []
        target_software_names_values = []
        target_software_main_classes_values = []
        target_software_additional_classes_values = []
        software_usage_status_values = []
        ru_software_availability_status_values = []
        ftr_values = []
        priority_values = []
        software_in_use_names_values = []
        reason_values = []
        comment_values = []

        companies_with_values = values.count(True)

        for company_index, company_id in enumerate(companies_ids):
            if not values[company_index]:
                continue

            company_name_values.append(companies[company_index].short_name)

            details = company_details_map.get(company_id, {}).get(bf_id, {})

            software_ids = details.get('software_ids', [])

            software_in_use_names = []

            for index_number, software_id in enumerate(software_ids, start=1):
                software_info = software_dict[software_id]

                if len(software_ids) > 1:
                    if software_name := software_info.get('name', ''):
                        software_in_use_names.append(f'\n-{index_number}- {software_name}')
                else:
                    if software_name := software_info.get('name', ''):
                        software_in_use_names.append(software_name)

            if target_software_id := details.get('target_software_id', None):
                target_software_info = software_dict[target_software_id]

                target_software_names_values.append(target_software_info.get('name', ''))
                target_software_main_classes_values.append(target_software_info.get('main_class_name', ''))
                target_software_additional_classes_values.append(target_software_info.get('additional_classes_names', ''))

            software_in_use_names_values.append('\n'.join(software_in_use_names).strip())

            software_usage_status_values.append(details.get('software_usage_status', ''))
            ru_software_availability_status_values.append(details.get('ru_software_availability_status', ''))
            ftr_values.append(details.get('ftr', ''))
            priority_values.append(details.get('priority', ''))
            reason_values.append(details.get('reason', ''))
            comment_values.append(details.get('comment', ''))

        for cell_values, column_index, column_width_param_name in zip(
                [
                    company_name_values, ftr_values, priority_values, reason_values,
                    ru_software_availability_status_values, software_usage_status_values,
                    target_software_names_values, target_software_main_classes_values,
                    target_software_additional_classes_values, software_in_use_names_values, comment_values,
                ],
                [5, 7, 9, 11, 12, 13, 14, 15, 16, 17, 18],
                [
                    'company_name', 'ftr', 'priority', 'reason', 'ru_software_availability_status',
                    'software_usage_status', 'target_software_names', 'target_software_main_classes',
                    'target_software_additional_classes', 'software_in_use_names', 'comment',
                ],
        ):
            if companies_with_values > 1:
                for company_index, cell_value in enumerate(cell_values):
                    cell_values[company_index] = f'{company_index + 1}) {cell_value}' if cell_value else ''

            cell_value = '\n\n'.join([cell_value for cell_value in cell_values if cell_value]).strip()
            insert_value_into_cell(row_pos, column_index, cell_value)
            if column_width_param_name:
                if len(cell_value) > 100:
                    new_width = len(cell_value) // 10 + 10
                    if new_width > 100:
                        new_width = 100
                    if new_width > columns_width[column_width_param_name]:
                        columns_width[column_width_param_name] = new_width

    def write_business_function_to_excel(bf, p, parent_bf_names):
        bf_id = bf.get('id', 0)
        level = bf.get('level', 0)
        name = bf.get('name', '')
        ftr_example = bf.get('ftr_example', '')
        priority_example = bf.get('priority_example', '')
        reason_example = bf.get('reason_example', '')
        children = bf.get('children', [])

        values = []

        for company_id in companies_ids:
            details = company_details_map.get(company_id, {}).get(bf_id, {})

            values.append(details.get('value') or False)

        if True not in values and level == 3:
            return p

        if 0 not in parent_bf_names and level == 1:
            parent_bf_names[0] = ''

        if parent_bf_names and list(parent_bf_names.keys())[-1] > level:
            del parent_bf_names[list(parent_bf_names.keys())[-1]]

        parent_bf_names[level] = name

        if children:
            if level == 2 and True in values:
                for col_pos, name in enumerate(parent_bf_names.values(), start=1):
                    insert_value_into_cell(p, col_pos, name)
                insert_value_into_cell(p, 6, ftr_example)
                insert_value_into_cell(p, 8, priority_example)
                insert_value_into_cell(p, 10, reason_example)
                insert_row_data(p, bf_id, values)
                p += 1

            for child in children:
                p = write_business_function_to_excel(child, p, parent_bf_names)

        if True in values and level == 3:
            for col_pos, name in enumerate(parent_bf_names.values(), start=1):
                insert_value_into_cell(p, col_pos, name)
            insert_value_into_cell(p, 6, ftr_example)
            insert_value_into_cell(p, 8, priority_example)
            insert_value_into_cell(p, 10, reason_example)
            insert_row_data(p, bf_id, values)
            return p + 1

        return p

    ws.freeze_panes = 'A2'

    ws.column_dimensions['A'].width = columns_width['bf_name']
    ws.column_dimensions['B'].width = columns_width['bf_name']
    ws.column_dimensions['C'].width = columns_width['bf_name']
    ws.column_dimensions['D'].width = columns_width['bf_name']
    ws.column_dimensions['E'].width = columns_width['company_name']
    ws.column_dimensions['F'].width = columns_width['bf_ftr_example']
    ws.column_dimensions['G'].width = columns_width['ftr']
    ws.column_dimensions['H'].width = columns_width['bf_priority_example']
    ws.column_dimensions['I'].width = columns_width['priority']
    ws.column_dimensions['J'].width = columns_width['bf_reason_example']
    ws.column_dimensions['K'].width = columns_width['reason']
    ws.column_dimensions['L'].width = columns_width['ru_software_availability_status']
    ws.column_dimensions['M'].width = columns_width['software_usage_status']
    ws.column_dimensions['N'].width = columns_width['target_software_names']
    ws.column_dimensions['O'].width = columns_width['target_software_main_classes']
    ws.column_dimensions['P'].width = columns_width['target_software_additional_classes']
    ws.column_dimensions['Q'].width = columns_width['software_in_use_names']
    ws.column_dimensions['R'].width = columns_width['comment']

    insert_value_into_cell(1, 1, 'Бизнес-функции', True)
    insert_value_into_cell(1, 2, 'Бизнес-функции 1 уровня', True)
    insert_value_into_cell(1, 3, 'Бизнес-функции 2 уровня', True)
    insert_value_into_cell(1, 4, 'Бизнес-функции 3 уровня', True)
    insert_value_into_cell(1, 5, 'Предприятия', True)
    insert_value_into_cell(1, 6, 'Функционально-технические требования (образец)', True)
    insert_value_into_cell(1, 7, 'Функционально-технические требования', True)
    insert_value_into_cell(1, 8, 'Приоритет (образец)', True)
    insert_value_into_cell(1, 9, 'Приоритет', True)
    insert_value_into_cell(1, 10, 'Обоснование выбора приоритета (образец)', True)
    insert_value_into_cell(1, 11, 'Обоснование выбора приоритета', True)
    insert_value_into_cell(1, 12, 'Оценка возможностей использования российских продуктов', True)
    insert_value_into_cell(1, 13, 'Статус использования ПО', True)
    insert_value_into_cell(1, 14, 'Целевое ПО', True)
    insert_value_into_cell(1, 15, 'Основной класс целевого ПО', True)
    insert_value_into_cell(1, 16, 'Доп классы целевого ПО', True)
    insert_value_into_cell(1, 17, 'Используемое ПО', True)
    insert_value_into_cell(1, 18, 'Комментарии', True)

    for root_business_function in root_business_functions:
        business_functions_tree = build_tree(root_business_function)

        position = write_business_function_to_excel(business_functions_tree, position, {})

    def merge_cells_in_column(column_index):
        current_value_rows = []
        row_counter = 2
        prev_cell_value = None
        while True:
            cell = ws.cell(row=row_counter, column=column_index)

            if prev_cell_value:
                if prev_cell_value == cell.value:
                    current_value_rows.append(row_counter)
                else:
                    ws.merge_cells(
                        start_row=current_value_rows[0],
                        end_row=current_value_rows[-1],
                        start_column=column_index,
                        end_column=column_index,
                    )
                    current_value_rows = [row_counter]
                    prev_cell_value = cell.value
            else:
                prev_cell_value = cell.value
                current_value_rows.append(row_counter)

            if not cell.value:
                break

            row_counter += 1

    merge_cells_in_column(1)
    merge_cells_in_column(2)
    merge_cells_in_column(3)

    if not consolidated:
        ws.delete_cols(5)
        ws.column_dimensions['E'].width = columns_width['bf_ftr_example']
        ws.column_dimensions['F'].width = columns_width['ftr']
        ws.column_dimensions['G'].width = columns_width['bf_priority_example']
        ws.column_dimensions['H'].width = columns_width['priority']
        ws.column_dimensions['I'].width = columns_width['bf_reason_example']
        ws.column_dimensions['J'].width = columns_width['reason']
        ws.column_dimensions['K'].width = columns_width['ru_software_availability_status']
        ws.column_dimensions['L'].width = columns_width['software_usage_status']
        ws.column_dimensions['M'].width = columns_width['target_software_names']
        ws.column_dimensions['N'].width = columns_width['target_software_main_classes']
        ws.column_dimensions['O'].width = columns_width['target_software_additional_classes']
        ws.column_dimensions['P'].width = columns_width['software_in_use_names']
        ws.column_dimensions['Q'].width = columns_width['comment']
    else:
        ws.column_dimensions['E'].width = columns_width['company_name']
        ws.column_dimensions['I'].width = columns_width['priority']
        ws.column_dimensions['L'].width = columns_width['ru_software_availability_status']
        ws.column_dimensions['M'].width = columns_width['software_usage_status']
        ws.column_dimensions['G'].width = columns_width['ftr']
        ws.column_dimensions['K'].width = columns_width['reason']
        ws.column_dimensions['N'].width = columns_width['target_software_names']
        ws.column_dimensions['O'].width = columns_width['target_software_main_classes']
        ws.column_dimensions['P'].width = columns_width['target_software_additional_classes']
        ws.column_dimensions['Q'].width = columns_width['software_in_use_names']
        ws.column_dimensions['R'].width = columns_width['comment']

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    return FileResponse(
        buffer,
        filename=(
            f'БизнесФункции_{"Консолидировано_" if consolidated else ""}'
            f'{user_company.short_name.replace(" ", "").replace("»", "_").replace("«", "_")}.xlsx'
        ),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        status=HTTPStatus.OK,
        as_attachment=True,
    )
